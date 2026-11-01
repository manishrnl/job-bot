"""SQLite job history with a formatted Excel export, so applications are trackable and readable."""
import csv
import datetime
import sqlite3

from pathlib import Path

from .config import DATA

COLS = ["platform", "job_id", "title", "company", "url", "status", "score", "fit", "resume", "note", "at"]
RUNS_DIR = DATA / "runs"
RUNS_DIR.mkdir(exist_ok=True)
_db = sqlite3.connect(DATA / "jobs.db")
_db.execute("""CREATE TABLE IF NOT EXISTS jobs (
  platform TEXT, job_id TEXT, title TEXT, company TEXT, url TEXT, status TEXT,
  score INTEGER, resume TEXT, note TEXT, at TEXT, PRIMARY KEY (platform, job_id))""")
if "fit" not in [r[1] for r in _db.execute("PRAGMA table_info(jobs)")]:
    _db.execute("ALTER TABLE jobs ADD COLUMN fit INTEGER")

# Statuses that mean "don't look at this job again". Errors/stuck are retried on the next run.
FINAL = ("Applied", "Low-Match", "Low-Fit", "External", "Excluded", "Draft-Saved", "Draft-Local")


def seen(platform, job_id, statuses=FINAL):
    q = f"SELECT 1 FROM jobs WHERE platform=? AND job_id=? AND status IN ({','.join('?' * len(statuses))})"
    return _db.execute(q, (platform, job_id, *statuses)).fetchone() is not None


def record(platform, job_id, title="", company="", url="", status="", score=None, resume="", note="", fit=None):
    row = dict(platform=platform, job_id=job_id, title=title, company=company, url=url, status=status,
               score=score, fit=fit, resume=str(resume or ""), note=note,
               at=datetime.datetime.now().isoformat(timespec="seconds"))
    _db.execute(f"INSERT OR REPLACE INTO jobs ({','.join(COLS)}) VALUES ({','.join('?' * len(COLS))})",
                [row[c] for c in COLS])
    _db.commit()
    export_excel()


def rows(limit=50, status=None):
    q = "SELECT at, platform, status, score, fit, title, company, url FROM jobs"
    args = []
    if status:
        q += " WHERE status=?"
        args.append(status)
    return _db.execute(q + " ORDER BY at DESC LIMIT ?", (*args, limit)).fetchall()


def stats():
    today = datetime.date.today().isoformat()
    by_status = dict(_db.execute("SELECT status, COUNT(*) FROM jobs GROUP BY status").fetchall())
    applied_today = _db.execute("SELECT COUNT(*) FROM jobs WHERE status='Applied' AND at >= ?", (today,)).fetchone()[0]
    avg = _db.execute("SELECT ROUND(AVG(score)) FROM jobs WHERE status='Applied'").fetchone()[0]
    per_day = _db.execute("SELECT substr(at,1,10) d, COUNT(*) FROM jobs WHERE status='Applied' "
                          "GROUP BY d ORDER BY d DESC LIMIT 14").fetchall()
    by_platform = dict(_db.execute("SELECT platform, COUNT(*) FROM jobs WHERE status='Applied' GROUP BY platform").fetchall())
    return dict(by_status=by_status, applied_today=applied_today, avg_score=avg, per_day=per_day,
                by_platform=by_platform, total_applied=by_status.get("Applied", 0))


def export_excel():
    """Rebuild the readable application workbook from SQLite history."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.worksheet.table import Table, TableStyleInfo

    path = DATA / "applications.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"
    band_fill = PatternFill("solid", fgColor="EAF2F8")
    status_cell_styles = {
        "Applied": ("E2F0D9", "006100"),
        "Error": ("FCE4D6", "9C0006"),
        "Stuck": ("FCE4D6", "9C0006"),
        "Draft-Saved": ("DDEBF7", "1F4E78"),
        "Draft-Local": ("DDEBF7", "1F4E78"),
        "Dry-Run": ("DDEBF7", "1F4E78"),
        "Low-Fit": ("FFF2CC", "9C6500"),
        "Low-Match": ("FFF2CC", "9C6500"),
        "Excluded": ("E7E6E6", "666666"),
        "External": ("DDEBF7", "1F4E78"),
    }
    headers = ["Date", "Status", "ATS score", "Fit", "Job title", "Company",
               "LinkedIn job", "Resume PDF", "Note"]
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="17365D")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(
            left=Side(style="medium", color="FFFFFF") if cell.column > 1 else Side(),
            right=Side(style="medium", color="FFFFFF"),
            bottom=Side(style="medium", color="FFFFFF"),
        )

    rows = _db.execute(f"SELECT {','.join(COLS)} FROM jobs ORDER BY at DESC").fetchall()
    for row in rows:
        record = dict(zip(COLS, row))
        status_label = str(record["status"] or "").replace("_", "-").title()
        ws.append([record["at"], status_label, record["score"] or "-", record["fit"] or "-",
                   record["title"], record["company"], record["url"] or "N/A",
                   record["resume"] or "N/A", record["note"] or "N/A"])
        n = ws.max_row
        try:
            from datetime import datetime
            ws.cell(n, 1).value = datetime.fromisoformat(record["at"])
            ws.cell(n, 1).number_format = "ddd, mmm dd, yyyy, hh:mm"
            ws.cell(n, 1).font = Font(name="Calibri", size=11, bold=True, color="002060")
        except (TypeError, ValueError):
            pass
        title_cell = ws.cell(n, 5)
        title_cell.font = Font(name="Calibri", size=11, bold=True, color="C00000")
        title_cell.alignment = Alignment(vertical="top", wrap_text=True)

        job_cell = ws.cell(n, 7)
        if record["url"]:
            target = str(record["url"]).strip()
            if target.startswith("www."):
                target = "https://" + target
            if not target.lower().startswith(("https://", "http://")):
                target = "https://" + target
            job_cell.value = target
            link_target = f"microsoft-edge:{target}" if "linkedin.com/" in target.lower() else target
            job_cell.hyperlink = link_target
            job_cell.font = Font(name="Calibri", size=11, color="0563C1", underline="single")

        resume_cell = ws.cell(n, 8)
        resume_path = Path(record["resume"]) if record["resume"] else None
        if resume_path:
            resume_cell.hyperlink = resume_path.resolve().as_uri()
            resume_cell.font = Font(name="Calibri", size=11, color="008000", underline="single")

        for col in (2, 3, 4, 6, 7, 8):
            ws.cell(n, col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=col in (7, 8))
        ws.cell(n, 5).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.cell(n, 9).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col in range(1, 10):
            cell = ws.cell(n, col)
            if n % 2 == 1:
                cell.fill = band_fill
            cell.border = Border(
                left=Side(style="medium", color="FFFFFF") if col > 1 else Side(),
                right=Side(style="medium", color="FFFFFF"),
                top=Side(style="medium", color="FFFFFF"),
                bottom=Side(style="medium", color="FFFFFF"),
            )
            if cell.value in ("-", "N/A"):
                cell.font = Font(name="Calibri", size=11, bold=True, color="002060")
        # Apply score colors directly so they display in viewers that do not
        # evaluate conditional formatting; the rules below cover future edits.
        score_colors = (
            (90, "E2F0D9", "006100"),
            (80, "C6E0B4", "375623"),
            (70, "FFF2CC", "9C6500"),
            (None, "FCE4D6", "9C0006"),
        )
        for col in (3, 4):
            cell = ws.cell(n, col)
            score = cell.value
            if isinstance(score, (int, float)):
                if score > 90:
                    fill_color, font_color = score_colors[0][1:]
                elif score >= 80:
                    fill_color, font_color = score_colors[1][1:]
                elif score >= 70:
                    fill_color, font_color = score_colors[2][1:]
                else:
                    fill_color, font_color = score_colors[3][1:]
                cell.fill = PatternFill("solid", fgColor="FF" + fill_color)
                cell.font = Font(name="Calibri", size=11, bold=True, color="FF" + font_color)
        status_cell = ws.cell(n, 2)
        if status_cell.value in status_cell_styles:
            fill_color, font_color = status_cell_styles[status_cell.value]
            status_cell.fill = PatternFill("solid", fgColor="FF" + fill_color)
            status_cell.font = Font(name="Calibri", size=11, bold=True, color="FF" + font_color)
        ws.row_dimensions[n].height = 24

    last_row = ws.max_row
    if last_row >= 2:
        # Color ATS and fit scores using the same four score bands.
        score_styles = [
            (">90", "E2F0D9", "006100"),
            ("80-90", "C6E0B4", "375623"),
            ("70-80", "FFF2CC", "9C6500"),
            ("<70", "FCE4D6", "9C0006"),
        ]
        for column in ("C", "D"):
            score_range = f"{column}2:{column}{last_row}"
            for band, fill_color, font_color in score_styles:
                if band == ">90":
                    formula = [f"AND(ISNUMBER({column}2),{column}2>90)"]
                elif band == "80-90":
                    formula = [f"AND(ISNUMBER({column}2),{column}2>=80,{column}2<=90)"]
                elif band == "70-80":
                    formula = [f"AND(ISNUMBER({column}2),{column}2>=70,{column}2<80)"]
                else:
                    formula = [f"AND(ISNUMBER({column}2),{column}2<70)"]
                ws.conditional_formatting.add(
                    score_range,
                    FormulaRule(formula=formula,
                                fill=PatternFill("solid", fgColor=fill_color),
                                font=Font(color=font_color, bold=True)),
                )

        # Color by status so the state is easy to scan down the sheet.
        status_styles = {
            "applied": ("E2F0D9", "006100"),
            "error": ("FCE4D6", "9C0006"),
            "stuck": ("FCE4D6", "9C0006"),
            "draft_saved": ("DDEBF7", "1F4E78"),
            "draft_local": ("DDEBF7", "1F4E78"),
            "dry_run": ("DDEBF7", "1F4E78"),
            "low_fit": ("FFF2CC", "9C6500"),
            "low_match": ("FFF2CC", "9C6500"),
            "excluded": ("E7E6E6", "666666"),
            "external": ("DDEBF7", "1F4E78"),
        }
        for status, (fill_color, font_color) in status_styles.items():
            ws.conditional_formatting.add(
                f"B2:B{last_row}",
                FormulaRule(
                    formula=[f'$B2="{status.replace("_", "-").title()}"'],
                    fill=PatternFill("solid", fgColor=fill_color),
                    font=Font(color=font_color, bold=True),
                ),
            )

    ws.freeze_panes = "B2"
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 26
    for column, width in {"A": 19, "B": 16, "C": 11, "D": 10, "E": 34,
                          "F": 26, "G": 44, "H": 44.55, "I": 30.22}.items():
        ws.column_dimensions[column].width = width
    table = Table(displayName="Table1", ref=f"A1:I{ws.max_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight2", showRowStripes=True,
                                          showColumnStripes=False)
    ws.add_table(table)
    try:
        wb.save(path)
    except PermissionError:
        # Excel locks the workbook while it is open. Keep the tracker run
        # working and provide a refreshed copy the user can open immediately.
        path = DATA / "applications_updated.xlsx"
        wb.save(path)
    return path


def write_run_summary(platform, submitted=0, dry_run=False, limit=None, counts=None):
    """Write one timestamped CSV summary per run so you can measure real submissions."""
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    path = RUNS_DIR / f"{platform}_{stamp}_summary.csv"
    counts = counts or {}
    row = [
        "run_at",
        "mode",
        "limit",
        "submitted",
        "applied",
        "dry_run",
        "low_fit",
        "low_match",
        "error",
    ]
    data = [
        datetime.datetime.now().isoformat(timespec="seconds"),
        "Dry-Run" if dry_run else "live",
        str(limit or ""),
        str(submitted),
        str(counts.get("Applied", 0)),
        str(counts.get("Dry-Run", 0)),
        str(counts.get("Low-Fit", 0)),
        str(counts.get("Low-Match", 0)),
        str(counts.get("Error", 0)),
    ]
    with open(path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(row)
        w.writerow(data)
    return path
